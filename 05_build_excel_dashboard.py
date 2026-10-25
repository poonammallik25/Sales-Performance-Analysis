"""Builds 03_Dashboard/Sales_Dashboard.xlsx (formula-driven Excel dashboard). Run after 01_clean_and_analyze.py."""
import os
import pandas as pd
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

df = pd.read_csv(os.path.join(ROOT, '01_Data', 'sales_clean.csv'), keep_default_na=False)
N = len(df); last = N + 1
NAVY, TEAL, ORANGE, RED, LIGHT = '1F3A5F', '2A9D8F', 'E9A23B', 'D1495B', 'EEF2F7'
F = lambda **k: Font(name='Arial', **k)
wb = Workbook()
dash = wb.active; dash.title = 'Dashboard'
st = wb.create_sheet('Summary_Tables'); data = wb.create_sheet('Clean_Data'); notes = wb.create_sheet('Data_Notes')

# ---------------- Clean_Data
cols = list(df.columns)
data.append(cols)
for r in df.itertuples(index=False):
    row = list(r); di = cols.index('Date'); row[di] = pd.to_datetime(row[di]).to_pydatetime(); data.append(row)
for c in data[1]: c.font = F(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY); c.alignment = Alignment(horizontal='center')
for i, c in enumerate(cols, 1):
    data.column_dimensions[L(i)].width = max(12, len(c) + 3)
    if c == 'Date':
        for x in data[L(i)][1:]: x.number_format = 'yyyy-mm-dd'
    if c in ('Gross_Sales', 'Discounts', 'Sales', 'COGS', 'Profit'):
        for x in data[L(i)][1:]: x.number_format = '#,##0.00'
data.freeze_panes = 'A2'
t = Table(displayName='SalesData', ref=f'A1:{L(len(cols))}{last}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); data.add_table(t)
col = {c: L(i + 1) for i, c in enumerate(cols)}
rng = lambda c: f"Clean_Data!${col[c]}$2:${col[c]}${last}"

# ---------------- Summary_Tables
st.sheet_view.showGridLines = False
for i, w in enumerate([22, 16, 16, 16, 16, 16, 16], 1): st.column_dimensions[L(i)].width = w
st['A1'] = 'Summary Tables (all values are live formulas on Clean_Data)'; st['A1'].font = F(bold=True, size=14, color=NAVY)
def header(r, labels):
    for i, h in enumerate(labels, 1):
        c = st.cell(r, i, h); c.font = F(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY); c.alignment = Alignment(horizontal='center', wrap_text=True)
def title(r, txt): st.cell(r, 1, txt).font = F(bold=True, size=11, color=TEAL)
def fmt(r1, r2, spec):
    for r in range(r1, r2 + 1):
        for cc, nf in spec.items(): st[f'{cc}{r}'].number_format = nf; st[f'{cc}{r}'].font = F()
USD, PCT, INT = '$#,##0;[Red]-$#,##0', '0.0%;[Red]-0.0%', '#,##0'

def dim_table(r0, name, dimcol, items):
    title(r0, name); header(r0 + 1, [dimcol.replace('_', ' '), 'Net Sales', 'Profit', 'Margin %', 'Units Sold', 'Profit Share %', 'Orders'])
    a = r0 + 2
    for i, it in enumerate(items):
        r = a + i; st[f'A{r}'] = it; st[f'A{r}'].font = F()
        st[f'B{r}'] = f'=SUMIFS({rng("Sales")},{rng(dimcol)},$A{r})'
        st[f'C{r}'] = f'=SUMIFS({rng("Profit")},{rng(dimcol)},$A{r})'
        st[f'D{r}'] = f'=IFERROR(C{r}/B{r},0)'
        st[f'E{r}'] = f'=SUMIFS({rng("Units_Sold")},{rng(dimcol)},$A{r})'
        st[f'F{r}'] = f'=C{r}/C${a+len(items)}'
        st[f'G{r}'] = f'=COUNTIFS({rng(dimcol)},$A{r})'
    z = a + len(items); st[f'A{z}'] = 'Total'
    for cc in 'BCEG': st[f'{cc}{z}'] = f'=SUM({cc}{a}:{cc}{z-1})'
    st[f'D{z}'] = f'=C{z}/B{z}'; st[f'F{z}'] = f'=SUM(F{a}:F{z-1})'
    fmt(a, z, {'B': USD, 'C': USD, 'D': PCT, 'E': INT, 'F': PCT, 'G': INT})
    for cc in 'ABCDEFG': st[f'{cc}{z}'].font = F(bold=True); st[f'{cc}{z}'].fill = PatternFill('solid', fgColor=LIGHT)
    return a, z - 1

segs = ['Government', 'Small Business', 'Enterprise', 'Midmarket', 'Channel Partners']
ctrs = ['France', 'Germany', 'Canada', 'United States of America', 'Mexico']
prods = ['Paseo', 'VTT', 'Amarilla', 'Velo', 'Montana', 'Carretera']
s1, s2 = dim_table(3, '1. Performance by Segment', 'Segment', segs)           # rows 5-9, total 10
c1, c2 = dim_table(13, '2. Performance by Country', 'Country', ctrs)          # 15-19, total 20
p1, p2 = dim_table(23, '3. Performance by Product', 'Product', prods)         # 25-30, total 31

# discount bands
r0 = 34; title(r0, '4. Discount Band Analysis'); header(r0 + 1, ['Discount Band', 'Orders', 'Gross Sales', 'Discounts', 'Avg Discount %', 'Profit', 'Margin %'])
bands = ['No Discount', 'Low', 'Medium', 'High']; d1 = r0 + 2
for i, b in enumerate(bands):
    r = d1 + i; st[f'A{r}'] = b
    st[f'B{r}'] = f'=COUNTIFS({rng("Discount_Band")},$A{r})'
    st[f'C{r}'] = f'=SUMIFS({rng("Gross_Sales")},{rng("Discount_Band")},$A{r})'
    st[f'D{r}'] = f'=SUMIFS({rng("Discounts")},{rng("Discount_Band")},$A{r})'
    st[f'E{r}'] = f'=IFERROR(D{r}/C{r},0)'
    st[f'F{r}'] = f'=SUMIFS({rng("Profit")},{rng("Discount_Band")},$A{r})'
    st[f'G{r}'] = f'=IFERROR(F{r}/SUMIFS({rng("Sales")},{rng("Discount_Band")},$A{r}),0)'
d2 = d1 + 3; fmt(d1, d2, {'B': INT, 'C': USD, 'D': USD, 'E': PCT, 'F': USD, 'G': PCT})
for r in range(d1, d2 + 1): st[f'A{r}'].font = F()

# monthly
r0 = 42; title(r0, '5. Monthly Trend'); header(r0 + 1, ['Year-Month', 'Net Sales', 'Profit', 'Margin %', 'MoM Sales Growth %'])
months = sorted(df.Year_Month.unique()); m1 = r0 + 2
for i, m in enumerate(months):
    r = m1 + i; st[f'A{r}'] = m
    st[f'B{r}'] = f'=SUMIFS({rng("Sales")},{rng("Year_Month")},$A{r})'
    st[f'C{r}'] = f'=SUMIFS({rng("Profit")},{rng("Year_Month")},$A{r})'
    st[f'D{r}'] = f'=IFERROR(C{r}/B{r},0)'
    if i: st[f'E{r}'] = f'=B{r}/B{r-1}-1'
m2 = m1 + len(months) - 1; fmt(m1, m2, {'B': USD, 'C': USD, 'D': PCT, 'E': PCT})
for r in range(m1, m2 + 1): st[f'A{r}'].font = F()

# matrix
r0 = m2 + 3; title(r0, '6. Profit Margin % - Segment x Discount Band'); header(r0 + 1, ['Segment'] + bands); x1 = r0 + 2
for i, sg in enumerate(segs):
    r = x1 + i; st[f'A{r}'] = sg; st[f'A{r}'].font = F()
    for j, b in enumerate(bands):
        cc = L(j + 2)
        st[f'{cc}{r}'] = (f'=IFERROR(SUMIFS({rng("Profit")},{rng("Segment")},$A{r},{rng("Discount_Band")},{cc}${r0+1})'
                          f'/SUMIFS({rng("Sales")},{rng("Segment")},$A{r},{rng("Discount_Band")},{cc}${r0+1}),0)')
        st[f'{cc}{r}'].number_format = PCT; st[f'{cc}{r}'].font = F()
x2 = x1 + 4
st.conditional_formatting.add(f'B{x1}:E{x2}', CellIsRule(operator='lessThan', formula=['0'], fill=PatternFill('solid', bgColor='F8C9D0')))
st.conditional_formatting.add(f'B{x1}:E{x2}', CellIsRule(operator='greaterThan', formula=['0.2'], fill=PatternFill('solid', bgColor='CDEBE7')))

# like for like
r0 = x2 + 3; title(r0, '7. Like-for-Like Growth (Sep-Dec only; 2013 data starts in Sep)'); header(r0 + 1, ['Period', 'Net Sales', 'Profit'])
l1 = r0 + 2
for i, y in enumerate([2013, 2014]):
    r = l1 + i; st[f'A{r}'] = y
    st[f'B{r}'] = f'=SUMIFS({rng("Sales")},{rng("Year")},$A{r},{rng("Month_Number")},">=9")'
    st[f'C{r}'] = f'=SUMIFS({rng("Profit")},{rng("Year")},$A{r},{rng("Month_Number")},">=9")'
    st[f'A{r}'].number_format = '0'; st[f'A{r}'].font = F()
st[f'A{l1+2}'] = 'Growth %'; st[f'B{l1+2}'] = f'=B{l1+1}/B{l1}-1'; st[f'C{l1+2}'] = f'=C{l1+1}/C{l1}-1'
fmt(l1, l1 + 1, {'B': USD, 'C': USD}); fmt(l1 + 2, l1 + 2, {'B': PCT, 'C': PCT})
for cc in 'ABC': st[f'{cc}{l1+2}'].font = F(bold=True)

# enterprise what-if
r0 = l1 + 5; title(r0, '8. Enterprise What-If (upper bound: assumes same volume)'); header(r0 + 1, ['Metric', 'Value'])
e = r0 + 2
st[f'A{e}'] = 'Actual profit'; st[f'B{e}'] = f'=C{s1+2}'
st[f'A{e+1}'] = 'Discounts given'; st[f'B{e+1}'] = f'=SUMIFS({rng("Discounts")},{rng("Segment")},"Enterprise")'
st[f'A{e+2}'] = 'Profit with zero discounts'; st[f'B{e+2}'] = f'=B{e}+B{e+1}'
st[f'A{e+3}'] = 'Gross margin before discount'; st[f'B{e+3}'] = f'=(SUMIFS({rng("Gross_Sales")},{rng("Segment")},"Enterprise")-SUMIFS({rng("COGS")},{rng("Segment")},"Enterprise"))/SUMIFS({rng("Gross_Sales")},{rng("Segment")},"Enterprise")'
for i in range(3): st[f'B{e+i}'].number_format = USD
st[f'B{e+3}'].number_format = PCT
for i in range(4): st[f'A{e+i}'].font = F(); st[f'B{e+i}'].font = F()
st.freeze_panes = 'A3'

# ---------------- Dashboard
dash.sheet_view.showGridLines = False
for i in range(1, 15): dash.column_dimensions[L(i)].width = 13
dash.column_dimensions['A'].width = 2; dash.column_dimensions['N'].width = 2
dash.merge_cells('B1:M2'); dash['B1'] = 'SALES PERFORMANCE DASHBOARD  |  Sep 2013 - Dec 2014'
dash['B1'].font = F(bold=True, size=20, color='FFFFFF'); dash['B1'].alignment = Alignment(vertical='center', horizontal='left', indent=1)
for r in (1, 2):
    for c in range(2, 14): dash.cell(r, c).fill = PatternFill('solid', fgColor=NAVY)
dash.row_dimensions[1].height = 22; dash.row_dimensions[2].height = 22
kp = [('NET SALES', f'=SUM({rng("Sales")})', '$#,##0.0,,"M"', TEAL, 'After discounts'),
      ('PROFIT', f'=SUM({rng("Profit")})', '$#,##0.0,,"M"', TEAL, 'Sales - COGS'),
      ('PROFIT MARGIN', f'=SUM({rng("Profit")})/SUM({rng("Sales")})', '0.0%', ORANGE, 'Profit / Net Sales'),
      ('UNITS SOLD', f'=SUM({rng("Units_Sold")})', '#,##0.0,"K"', NAVY, 'All products'),
      ('DISCOUNTS GIVEN', f'=SUM({rng("Discounts")})', '$#,##0.0,,"M"', ORANGE, 'Gross - Net sales'),
      ('LOSS-MAKING ORDERS', f'=COUNTIFS({rng("Profit")},"<0")', '#,##0', RED, 'All in Enterprise')]
for i, (lab, f, nf, colr, sub) in enumerate(kp):
    c1_, c2_ = L(2 + 2 * i), L(3 + 2 * i)
    for r in (4, 5, 6):
        dash.merge_cells(f'{c1_}{r}:{c2_}{r}')
        for cc in (c1_, c2_): dash[f'{cc}{r}'].fill = PatternFill('solid', fgColor=LIGHT)
    dash[f'{c1_}4'] = lab; dash[f'{c1_}4'].font = F(bold=True, size=9, color='5B6B7C')
    dash[f'{c1_}5'] = f; dash[f'{c1_}5'].number_format = nf; dash[f'{c1_}5'].font = F(bold=True, size=20, color=colr)
    dash[f'{c1_}6'] = sub; dash[f'{c1_}6'].font = F(size=8, color='5B6B7C')
    for r in (4, 5, 6): dash[f'{c1_}{r}'].alignment = Alignment(horizontal='center', vertical='center')
dash.row_dimensions[5].height = 34
dash['B3'] = ''

def style(ch, ttl, w=15.5, h=7.6):
    ch.title = ttl; ch.width = w; ch.height = h; ch.legend.position = 'b'; return ch
# chart 1 segment
c = BarChart(); c.type = 'col'; style(c, 'Net Sales vs Profit by Segment')
c.add_data(Reference(st, min_col=2, max_col=3, min_row=s1 - 1, max_row=s2), titles_from_data=True); c.set_categories(Reference(st, min_col=1, min_row=s1, max_row=s2))
c.series[0].graphicalProperties.solidFill = NAVY; c.series[1].graphicalProperties.solidFill = TEAL; c.y_axis.numFmt = '$#,##0,,"M"'; c.y_axis.delete = False; c.x_axis.delete = False
dash.add_chart(c, 'B8')
# chart 2 monthly
c = LineChart(); style(c, 'Monthly Net Sales & Profit')
c.add_data(Reference(st, min_col=2, max_col=3, min_row=m1 - 1, max_row=m2), titles_from_data=True); c.set_categories(Reference(st, min_col=1, min_row=m1, max_row=m2))
c.series[0].graphicalProperties.line.solidFill = NAVY; c.series[1].graphicalProperties.line.solidFill = TEAL
for s in c.series: s.smooth = False; s.marker.symbol = 'circle'; s.marker.size = 5
c.y_axis.numFmt = '$#,##0,,"M"'; c.y_axis.delete = False; c.x_axis.delete = False; c.x_axis.tickLblPos = 'low'
dash.add_chart(c, 'H8')
# chart 3 discount band margin
c = BarChart(); c.type = 'col'; style(c, 'Profit Margin % by Discount Band'); c.legend = None
c.add_data(Reference(st, min_col=7, min_row=d1 - 1, max_row=d2), titles_from_data=True); c.set_categories(Reference(st, min_col=1, min_row=d1, max_row=d2))
c.series[0].graphicalProperties.solidFill = ORANGE; c.y_axis.numFmt = '0%'; c.dataLabels = DataLabelList(); c.dataLabels.showVal = True; c.dataLabels.showSerName = False; c.dataLabels.showCatName = False; c.dataLabels.showLegendKey = False; c.dataLabels.numFmt = '0.0%'; c.y_axis.delete = False; c.x_axis.delete = False
dash.add_chart(c, 'B25')
# chart 4 country profit
c = BarChart(); c.type = 'bar'; style(c, 'Profit by Country'); c.legend = None
c.add_data(Reference(st, min_col=3, min_row=c1 - 1, max_row=c2), titles_from_data=True); c.set_categories(Reference(st, min_col=1, min_row=c1, max_row=c2))
c.series[0].graphicalProperties.solidFill = NAVY; c.y_axis.numFmt = '$#,##0.0,,"M"'; c.x_axis.scaling.orientation = 'maxMin'; c.y_axis.crosses = 'max'; c.y_axis.delete = False; c.x_axis.delete = False
dash.add_chart(c, 'H25')

dash.merge_cells('B42:M42'); dash['B42'] = 'KEY INSIGHTS (data: Sep 2013 - Dec 2014)'; dash['B42'].font = F(bold=True, size=12, color='FFFFFF')
for cc in range(2, 14): dash.cell(42, cc).fill = PatternFill('solid', fgColor=TEAL)
ins = ['1. Government drives 67% of total profit ($11.4M at a 21.7% margin).',
       '2. Enterprise is the 3rd-largest segment ($19.6M sales) but LOSES money: -$0.6M (-3.1% margin).',
       '3. Root cause: Enterprise price $125 vs cost $120/unit = only 4% margin before any discount.',
       '4. Deeper discounts shrink margin in every segment: 21.9% (none) > 17.9% (low) > 14.4% (medium) > 9.1% (high).',
       '5. Like-for-like (Sep-Dec): sales +36.9% and profit +40.1% in 2014 vs 2013.',
       '6. Channel Partners earn a 73% margin but are only 1.5% of sales - a growth opportunity.',
       '7. Sales peak in Oct and Dec 2014 (~$12M each) - strong Q4 seasonality.']
for i, t_ in enumerate(ins):
    r = 43 + i; dash.merge_cells(f'B{r}:M{r}'); dash[f'B{r}'] = t_; dash[f'B{r}'].font = F(size=10); dash[f'B{r}'].alignment = Alignment(indent=1, vertical='center')
    dash.row_dimensions[r].height = 18

# ---------------- Data notes
notes.sheet_view.showGridLines = False; notes.column_dimensions['A'].width = 6; notes.column_dimensions['B'].width = 38; notes.column_dimensions['C'].width = 85
notes['A1'] = 'Data Cleaning Log & Dictionary'; notes['A1'].font = F(bold=True, size=14, color=NAVY)
rows = [('#', 'Issue / Field', 'Action / Meaning'),
        (1, "Column ' Sales' had a leading space", 'Header trimmed; spaces replaced with underscores'),
        (2, 'Discount Band blank in 53 rows', 'All 53 have Discounts = 0 -> relabelled "No Discount" (not deleted)'),
        (3, '36 rows with half units (e.g. 4492.5)', 'Kept as reported; flagged to data owner'),
        (4, 'Manufacturing Price does not reconcile to COGS / Units', 'COGS used as reported; flagged as a data-governance issue'),
        (5, '2013 has only Sep-Dec; 2014 is a full year', 'Growth measured like-for-like (Sep-Dec only)'),
        (6, 'Checks: Gross - Discounts = Sales; Sales - COGS = Profit', 'Validated on all 700 rows; 0 duplicates, 0 nulls after cleaning'),
        ('', '', ''), ('', 'Added columns', 'Profit_Margin_Pct, Discount_Pct, Year_Month, Quarter, Loss_Making, COGS_per_Unit'),
        ('', 'Net Sales', 'Gross Sales minus Discounts (column "Sales")'), ('', 'Profit', 'Net Sales minus COGS'), ('', 'Margin %', 'Profit divided by Net Sales')]
for i, rw in enumerate(rows, 3):
    for j, v in enumerate(rw, 1):
        c = notes.cell(i, j, v); c.font = F(bold=(i == 3), color='FFFFFF' if i == 3 else '000000'); c.alignment = Alignment(wrap_text=True, vertical='top')
        if i == 3: c.fill = PatternFill('solid', fgColor=NAVY)
from openpyxl.worksheet.properties import PageSetupProperties
dash.page_setup.orientation='landscape'; dash.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); dash.page_setup.fitToWidth=1; dash.page_setup.fitToHeight=1
dash.print_area='A1:N50'
wb.save(os.path.join(ROOT, '03_Dashboard', 'Sales_Dashboard.xlsx')); print('saved', s1, s2, c1, c2, p1, p2, d1, d2, m1, m2, x1, x2, l1, e)
